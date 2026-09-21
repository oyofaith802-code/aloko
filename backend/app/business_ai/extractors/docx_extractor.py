from pathlib import Path
from typing import Any

from docx import Document


def extract_docx(file_path: str | Path) -> dict[str, Any]:
    """
    Extract paragraphs and tables from a DOCX document.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {path}")

    document = Document(str(path))

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    tables = []

    for table_index, table in enumerate(document.tables):
        rows = []

        for row in table.rows:
            rows.append(
                [
                    cell.text.strip()
                    for cell in row.cells
                ]
            )

        tables.append(
            {
                "table_index": table_index,
                "rows": rows,
            }
        )

    full_text = "\n".join(paragraphs)

    return {
        "source_type": "docx",
        "filename": path.name,
        "paragraph_count": len(paragraphs),
        "paragraphs": paragraphs,
        "text": full_text,
        "tables": tables,
    }